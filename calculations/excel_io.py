"""Read the input tables and the reference results from the Excel workbook.

The workbook is the golden reference. Inputs are read as plain values; results are
read from the cached formula values so the engine can be compared against them.
All times are handled as Excel serial numbers (days since 1899-12-30), exactly as
Excel does, so comparisons such as READY_TIME <= PICK_CUTOFF give the same answer.
"""
from __future__ import annotations

import datetime as dt
from dataclasses import dataclass

import pandas as pd
from openpyxl import load_workbook

EXCEL_EPOCH = dt.datetime(1899, 12, 30)


def to_serial(value):
    """Excel serial number for a datetime / timedelta; any other value is returned unchanged."""
    if isinstance(value, dt.datetime):
        return (value - EXCEL_EPOCH).total_seconds() / 86400
    if isinstance(value, dt.timedelta):
        return value.total_seconds() / 86400
    return value


def serial_to_datetime(serial: float) -> dt.datetime:
    return EXCEL_EPOCH + dt.timedelta(days=serial)


def _table(ws, header_row: int) -> pd.DataFrame:
    rows = ws.iter_rows(min_row=header_row, values_only=True)
    header = [str(h).strip() if h is not None else f"col{i}" for i, h in enumerate(next(rows))]
    data = [r for r in rows if any(v is not None for v in r)]
    return pd.DataFrame(data, columns=header)


@dataclass
class Inputs:
    snapshots: pd.DataFrame        # SNAPSHOT_NO, T (serial)
    config: dict                   # ETC filters + truck parameters
    lm: pd.DataFrame               # LM005S1 (labour segments, one copy per snapshot)
    ti: pd.DataFrame               # TI102 (open lines, one copy per snapshot)
    tc: pd.DataFrame               # TI102C (completed lines)
    shipments: pd.DataFrame        # OB_SHIPMENT
    orders: pd.DataFrame           # OB_ORDER
    line_map: pd.DataFrame         # OB_LINE_MAP
    temps: pd.DataFrame            # TEMP_OPERATORS skills grid
    worktype_names: dict           # WORKTYPE -> description (M123 / M123T)


def load_inputs(path: str) -> Inputs:
    wb = load_workbook(path, read_only=True, data_only=True)
    cfg_ws = wb["CONFIG"]
    config = {cfg_ws.cell(r, 1).value: cfg_ws.cell(r, 2).value for r in range(4, 11)}
    snaps = [(cfg_ws.cell(r, 1).value, cfg_ws.cell(r, 2).value) for r in range(14, 32)]
    snapshots = pd.DataFrame([(n, to_serial(t)) for n, t in snaps if n is not None], columns=["SNAPSHOT_NO", "T"])
    tcfg = wb["TRUCK_CONFIG"]
    for r in range(6, 9):
        config[tcfg.cell(r, 1).value] = tcfg.cell(r, 2).value
    config["ANALYSIS_DATE"] = to_serial(config["ANALYSIS_DATE"])

    lm = _table(wb["LM005S1"], 2)
    ti = _table(wb["TI102"], 2)
    tc = _table(wb["TI102C"], 2)
    for df, cols in [(lm, ["SNAPSHOT_TIME", "TRANSACTIONDATE", "STARTDATETIME", "ENDDATETIME"]),
                     (ti, ["SNAPSHOT_TIME", "TRANSACTIONDATE", "CHANGEDATETIME"]),
                     (tc, ["TRANSACTIONDATE", "MOVED_TO_C_AT"])]:
        for c in cols:
            df[c] = df[c].map(to_serial)

    shipments = _table(wb["OB_SHIPMENT"], 5).iloc[:, :15]
    for c in ["ORIGINAL_DEPARTURE", "DEPARTURE_TIME", "DEPARTURE_CHANGED_AT"]:
        shipments[c] = shipments[c].map(to_serial)
    raw_orders = _table(wb["OB_ORDER"], 5)
    orders = raw_orders.iloc[:, :12].copy()
    orders.columns = ["WHORDERID", "SHIPMENTID", "ORIGINAL_SHIPMENTID", "MOVED_AT", "ROUTEID", "STOPID",
                      "RELEASE_TIME", "ASSIGNMENTS", "LINES", "WORK_TYPES", "DATA_TYPE", "OB_SCENARIO_ID"]
    orders["WHORDERID"] = orders.WHORDERID.astype(str)
    # real-day workbook: customer of each order (optional columns)
    orders["CUSTOMER"] = raw_orders["CUSTOMER (OWNER)"] if "CUSTOMER (OWNER)" in raw_orders else None
    orders["CUSTOMER_NAME"] = raw_orders["CUSTOMER_NAME"] if "CUSTOMER_NAME" in raw_orders else None
    for c in ["MOVED_AT", "RELEASE_TIME"]:
        orders[c] = orders[c].map(to_serial)
    line_map = _table(wb["OB_LINE_MAP"], 5)
    line_map["WHORDERID"] = line_map.WHORDERID.astype(str)
    line_map["CUSTOMER"] = line_map["CUSTOMER (OWNER)"] if "CUSTOMER (OWNER)" in line_map else None

    tws = wb["TEMP_OPERATORS"]
    temps = pd.DataFrame([[tws.cell(r, c).value for c in range(1, 8)] for r in range(6, 16)],
                         columns=["RESOURCENAME", 1, 2, 3, 4, 5, 6])

    ids = {r[0]: r[5] for r in wb["M123"].iter_rows(min_row=3, values_only=True) if r[0]}
    desc = {r[1]: r[3] for r in wb["M123T"].iter_rows(min_row=3, values_only=True) if r[1]}
    worktype_names = {wt: desc.get(mid) for mid, wt in ids.items()}
    wb.close()
    return Inputs(snapshots, config, lm, ti, tc, shipments, orders, line_map, temps, worktype_names)


def load_reference(path: str) -> dict:
    """Cached Excel results of ETC_CALC, TRUCK_WT_ETC, TRUCK_ETC and TRUCK_TEMP_NEED (for parity tests)."""
    wb = load_workbook(path, read_only=True, data_only=True)
    out = {name: _table(wb[name], 5).map(to_serial)
           for name in ["ETC_CALC", "TRUCK_WT_ETC", "TRUCK_ETC", "TRUCK_TEMP_NEED"]}
    wb.close()
    return out
