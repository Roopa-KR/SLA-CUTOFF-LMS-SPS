"""Create the explicit simulated operator cross-training dataset.

The source warehouse records do not contain a qualification matrix. These
deterministic planning assumptions are stored separately from actual LM005S1
assignments and TI102C completions.
"""
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill


PATH = Path(__file__).resolve().parents[1] / "data" / "workbook.xlsx"

# Backup-area familiarity chosen around adjacent warehouse processes. Every
# person already exists in LM005S1; this does not create operators.
ASSIGNMENTS = {
    "amack": [(2, "Cross-trained - working knowledge")],
    "cburney": [(2, "Cross-trained - experienced backup")],
    "dclements": [(1, "Cross-trained - working knowledge"), (2, "Cross-trained - working knowledge")],
    "ekim": [(2, "Cross-trained - experienced backup"), (3, "Cross-trained - working knowledge")],
    "esica": [(2, "Cross-trained - working knowledge")],
    "hgroff": [(2, "Cross-trained - working knowledge")],
    "jmccraw": [(2, "Cross-trained - experienced backup")],
    "lhall": [(1, "Cross-trained - working knowledge")],
    "mmathis": [(1, "Cross-trained - experienced backup")],
    "psam": [(2, "Cross-trained - experienced backup")],
    "Shullett": [(4, "Cross-trained - working knowledge")],
    "tedwards": [(3, "Cross-trained - working knowledge"), (5, "Cross-trained - working knowledge")],
}


def main() -> None:
    wb = openpyxl.load_workbook(PATH)
    if "OPERATOR_EXPERIENCE" in wb.sheetnames:
        del wb["OPERATOR_EXPERIENCE"]
    ws = wb.create_sheet("OPERATOR_EXPERIENCE", 6)
    ws["A1"] = "SIMULATED OPERATOR EXPERIENCE AND WORK-TYPE FAMILIARITY"
    ws["A1"].font = Font(bold=True, size=14, color="FFFFFF")
    ws["A1"].fill = PatternFill("solid", fgColor="7F6000")
    ws.merge_cells("A1:E1")
    ws["A2"] = ("Planning assumptions only. These rows are not verified training records and must not be read as "
                  "historical work. Actual assignments and productivity remain in LM005S1 and TI102C.")
    ws.merge_cells("A2:E2")
    ws["A3"] = ("Familiarity may establish planning eligibility. Estimated destination pace is calculated separately "
                  "as MIN(actual source pace, actual destination per-person rate).")
    ws.merge_cells("A3:E3")
    headers = ["RESOURCENAME", "WORKTYPE", "EXPERIENCE_LEVEL", "EVIDENCE_TYPE", "ASSUMPTION_NOTES"]
    for col, header in enumerate(headers, 1):
        cell = ws.cell(5, col, header)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="BF9000")
        cell.alignment = Alignment(wrap_text=True)
    row = 6
    for operator, assignments in ASSIGNMENTS.items():
        for worktype, level in assignments:
            values = [operator, worktype, level, "SIMULATED PLANNING ASSUMPTION",
                      "Synthetic cross-training profile for scenario analysis; not verified history."]
            for col, value in enumerate(values, 1):
                ws.cell(row, col, value)
            row += 1
    ws.freeze_panes = "A6"
    ws.auto_filter.ref = f"A5:E{row - 1}"
    for column, width in {"A": 18, "B": 12, "C": 38, "D": 32, "E": 80}.items():
        ws.column_dimensions[column].width = width
    wb.calculation.fullCalcOnLoad = True
    wb.calculation.forceFullCalc = True
    wb.save(PATH)


if __name__ == "__main__":
    main()
