"""Add one explicit synthesized PICK CUTOFF PASSED - LATE example.

SH0908-1072 has two genuinely open lines at snapshot 11 (18:45) and is fully
picked by snapshot 12 (19:00). Its outbound departure is already synthesized,
so this scenario changes that synthesized departure from 19:30 to 19:15. With
the configured 30-minute loading allowance, its pick cutoff becomes 18:45 and
the normal status rules produce exactly one late-cutoff snapshot.
"""
from __future__ import annotations

import datetime as dt
from pathlib import Path

import openpyxl


PATH = Path(__file__).resolve().parents[1] / "data" / "workbook.xlsx"
SHIPMENT_ID = "SH0908-1072"
DEPARTURE = dt.datetime(2026, 9, 8, 19, 15)
SCENARIO_ID = "T07_PICK_CUTOFF_PASSED_LATE"
NOTE = ("Real route 1072: 2 items, last item released 17:54. Synthesized status-coverage scenario: "
        "pick cutoff 18:45 and departure 19:15; snapshot 11 intentionally demonstrates open work at cutoff.")


def main() -> None:
    workbook = openpyxl.load_workbook(PATH)
    sheet = workbook["OB_SHIPMENT"]
    headers = {cell.value: cell.column for cell in sheet[5] if cell.value}
    row = next((row for row in range(6, sheet.max_row + 1)
                if sheet.cell(row, headers["SHIPMENTID"]).value == SHIPMENT_ID), None)
    if row is None:
        raise RuntimeError(f"{SHIPMENT_ID} was not found in OB_SHIPMENT")

    sheet.cell(row, headers["ORIGINAL_DEPARTURE"]).value = DEPARTURE
    sheet.cell(row, headers["DEPARTURE_TIME"]).value = DEPARTURE
    sheet.cell(row, headers["DEPARTURE_CHANGED_AT"]).value = None
    # PICK_CUTOFF_TIME (final) remains the workbook formula: departure - loading minutes.
    sheet.cell(row, headers["OB_SCENARIO_ID"]).value = SCENARIO_ID
    sheet.cell(row, headers["NOTE"]).value = NOTE

    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    workbook.calculation.calcMode = "auto"
    workbook.save(PATH)


if __name__ == "__main__":
    main()
