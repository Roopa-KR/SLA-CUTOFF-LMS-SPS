"""End-to-end invariants for the cleaned real-day workbook."""
import math
import datetime as dt

import openpyxl

from calculations.excel_io import serial_to_datetime

REMOVED_SHIPMENTS = {"SH0908-1001", "SH0908-1075", "SH0908-1085"}


def test_source_keys_and_relationships_are_complete(results):
    inp = results.inputs
    shipment_ids = set(inp.shipments.SHIPMENTID)
    order_ids = set(inp.orders.WHORDERID)
    line_ids = set(inp.line_map.TI102ID)

    assert len(inp.shipments) == len(shipment_ids) == 33
    assert len(inp.orders) == len(order_ids) == 421
    assert len(inp.line_map) == len(line_ids) == 1163
    assert set(inp.orders.SHIPMENTID) <= shipment_ids
    assert set(inp.orders.ORIGINAL_SHIPMENTID) <= shipment_ids
    assert set(inp.line_map.WHORDERID) == order_ids
    assert set(inp.tc.TI102ID) == line_ids
    assert set(inp.ti.TI102ID) <= line_ids
    assert not shipment_ids & REMOVED_SHIPMENTS


def test_warehouse_identifiers_are_consistent(results):
    inp = results.inputs
    assert inp.config["WHID_FILTER"] == 100
    assert set(inp.lm.WHID) == {100}
    assert set(inp.ti.WHID) == {100}
    assert set(inp.tc.WHID) == {100}

    workbook = openpyxl.load_workbook("data/workbook.xlsx", read_only=True, data_only=True)
    try:
        assert {workbook["M150"].cell(row, 2).value for row in range(3, 9)} == {100}
        assert {workbook["M123"].cell(row, 2).value for row in range(3, 9)} == {100}
        assert "TRUCK_REALLOCATION" in workbook.sheetnames
        assert "TRUCK_REALLOCATION_MOVES" in workbook.sheetnames
        assert not {"TRUCK_TEMP_NEED", "TRUCK_TEMP_LM005S1", "TEMP_OPERATORS"} & set(workbook.sheetnames)
    finally:
        workbook.close()


def test_synthesized_departures_follow_the_documented_rule(results):
    inp = results.inputs
    ownership = inp.line_map[["TI102ID", "WHORDERID"]].merge(
        inp.orders[["WHORDERID", "SHIPMENTID"]], on="WHORDERID", how="left")
    items = inp.tc[["TI102ID", "CREATEDATETIME"]].merge(ownership, on="TI102ID", how="inner")
    last_release = items.groupby("SHIPMENTID").CREATEDATETIME.max()
    lead = dt.timedelta(minutes=inp.config["ORDER_LEAD_MINUTES"])
    loading = dt.timedelta(minutes=inp.config["LOADING_MINUTES"])

    for shipment in inp.shipments.itertuples():
        unrounded_cutoff = last_release[shipment.SHIPMENTID].to_pydatetime() + lead
        midnight = unrounded_cutoff.replace(hour=0, minute=0, second=0, microsecond=0)
        elapsed_minutes = (unrounded_cutoff - midnight).total_seconds() / 60
        expected_cutoff = midnight + dt.timedelta(minutes=math.ceil(elapsed_minutes / 15) * 15)
        assert shipment.ORIGINAL_DEPARTURE == shipment.DEPARTURE_TIME
        expected_departure = expected_cutoff + loading
        if shipment.OB_SCENARIO_ID == "T07_PICK_CUTOFF_PASSED_LATE":
            # Explicitly labelled status-coverage exception: one quarter-hour earlier.
            expected_departure -= dt.timedelta(minutes=15)
            assert shipment.SHIPMENTID == "SH0908-1072"
        assert abs((serial_to_datetime(shipment.DEPARTURE_TIME) - expected_departure).total_seconds()) < 0.001


def test_no_pick_occurs_after_its_synthesized_departure(results):
    inp = results.inputs
    ownership = inp.line_map[["TI102ID", "WHORDERID"]].merge(
        inp.orders[["WHORDERID", "SHIPMENTID"]], on="WHORDERID", how="left")
    picked = inp.tc.merge(ownership, on="TI102ID", how="inner")
    departures = dict(zip(inp.shipments.SHIPMENTID, inp.shipments.DEPARTURE_TIME))

    violations = picked[picked.apply(lambda row: row.MOVED_TO_C_AT > departures[row.SHIPMENTID] + 1e-9, axis=1)]
    assert violations.empty


def test_replay_rows_obey_snapshot_lifecycle_and_quantity_rules(results):
    inp = results.inputs
    snapshot_times = dict(zip(results.snapshots.SNAPSHOT_NO, results.snapshots["T"]))

    for replay in (inp.ti, inp.lm):
        expected = replay.SNAPSHOT_NO.map(snapshot_times)
        assert (replay.SNAPSHOT_TIME - expected).abs().max() < 1e-9
        assert (replay.TRANSACTIONDATE <= replay.SNAPSHOT_TIME + 1e-9).all()

    assert not inp.ti.duplicated(["SNAPSHOT_NO", "TI102ID"]).any()
    assert not inp.lm.duplicated(["SNAPSHOT_NO", "LM005ID"]).any()
    assert (inp.ti.CHANGEDATETIME <= inp.ti.SNAPSHOT_TIME + 1e-9).all()
    assert inp.lm.loc[inp.lm.WIPIND == "Y", "ENDDATETIME"].isna().all()
    assert inp.lm.loc[inp.lm.WIPIND == "N", "ENDDATETIME"].notna().all()
    effective_end = inp.lm.ENDDATETIME.fillna(inp.lm.SNAPSHOT_TIME)
    assert (inp.lm.STARTDATETIME <= effective_end + 1e-9).all()
    assert (effective_end <= inp.lm.SNAPSHOT_TIME + 1e-9).all()
    elapsed_seconds = (effective_end - inp.lm.STARTDATETIME) * 86400
    assert (elapsed_seconds - inp.lm.TOTALTIMEINSEC).abs().max() < 0.001

    for items in (inp.ti, inp.tc):
        assert (items.ORDERQTY == items.TOTALQTY).all()
        accounted = items.OPENQTY + items.PROCESSEDQTY + items.SHORTQTY + items.CUTQTY
        assert (items.TOTALQTY == accounted).all()
        assert (items[["ORDERQTY", "OPENQTY", "PROCESSEDQTY", "SHORTQTY", "CUTQTY"]] >= 0).all().all()

    assert set(inp.ti.CURRENTSTATUSID) <= {0, 2}
    assert set(inp.tc.CURRENTSTATUSID) == {3}
    assert (inp.tc.OPENQTY == 0).all()
    assert inp.tc.RESOURCENAME.fillna("").ne("").all()
    assert inp.ti.loc[inp.ti.CURRENTSTATUSID == 0, "RESOURCENAME"].fillna("").eq("").all()
    assert inp.ti.loc[inp.ti.CURRENTSTATUSID == 2, "RESOURCENAME"].fillna("").ne("").all()


def test_item_worktype_and_assignment_dependencies_are_stable(results):
    inp = results.inputs
    expected = inp.line_map.set_index("TI102ID")[["WORKTYPE", "ASSIGNMENTID"]]

    for items in (inp.ti, inp.tc):
        joined = items[["TI102ID", "WORKTYPE", "ASSIGNMENTID"]].join(
            expected, on="TI102ID", rsuffix="_SOURCE")
        assert (joined.WORKTYPE == joined.WORKTYPE_SOURCE).all()
        assert (joined.ASSIGNMENTID == joined.ASSIGNMENTID_SOURCE).all()


def test_workbook_checks_and_results_are_clean(results):
    workbook = openpyxl.load_workbook("data/workbook.xlsx", read_only=True, data_only=True)
    try:
        checks = workbook["TRUCK_CHECKS"]
        assert all(checks.cell(row, 6).value in {"PASS", "INFO"} for row in range(6, 26))
        truck_sheet = workbook["TRUCK_ETC"]
        assert truck_sheet["AI5"].value == "LINES_LEFT_AT_DEPARTURE"
        assert all((truck_sheet.cell(row, 35).value or 0) == 0 for row in range(6, truck_sheet.max_row + 1))
    finally:
        workbook.close()

    late = results.trucks[results.trucks.TRUCK_STATUS == "PICK CUTOFF PASSED - LATE"]
    assert late[["SNAPSHOT_NO", "SHIPMENTID", "OPEN_LINES"]].to_dict("records") == [
        {"SNAPSHOT_NO": 11, "SHIPMENTID": "SH0908-1072", "OPEN_LINES": 2}
    ]
    assert not results.trucks.TRUCK_STATUS.str.startswith("DEPARTED -").any()
    assert len(results.trucks) == 18 * 33


def test_simulated_familiarity_is_separate_and_headcount_is_conserved(results):
    completed_worktypes = results.inputs.tc.groupby("RESOURCENAME").WORKTYPE.nunique()
    assert (completed_worktypes <= 1).all()
    experience = results.inputs.operator_experience
    assert len(experience) > 0
    assert set(experience.EVIDENCE_TYPE) == {"SIMULATED PLANNING ASSUMPTION"}
    assert set(experience.RESOURCENAME) <= set(results.inputs.lm.RESOURCENAME)
    assert len(results.moves) == 3
    assert len(results.reallocation) == 13
    assert (results.reallocation.ALLOCATED_OPERATORS >= 0).all()
    assert (results.reallocation.REMAINING_SHORTFALL >= 0).all()
    assert (results.reallocation.PLANNING_STATUS_AFTER == "AT RISK").any()
    assert results.moves.FAMILIARITY_TYPE.eq("SIMULATED PLANNING ASSUMPTION").all()
    assert results.moves.SOURCE_STATUS_AFTER.eq("ON TIME").all()
    unique_wt = results.worktypes.drop_duplicates(["SNAPSHOT_NO", "WORKTYPE"])
    assert (unique_wt.groupby("SNAPSHOT_NO").RESOURCES.sum()
            == unique_wt.groupby("SNAPSHOT_NO").RESOURCES_AFTER_REALLOCATION.sum()).all()
